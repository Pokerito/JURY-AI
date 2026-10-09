import os
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "demo_contracts")
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ─────────────────────────────────────────────────────────────────────────────
# CONTRACT 11: AI MODEL TRAINING DATA LICENSING AGREEMENT
# ─────────────────────────────────────────────────────────────────────────────
c11_text = """AI FOUNDATION MODEL TRAINING DATA & CONTENT REPOSITORY LICENSE AGREEMENT

This AI FOUNDATION MODEL TRAINING DATA & CONTENT REPOSITORY LICENSE AGREEMENT (the "Agreement") is entered into as of November 14, 2024 (the "Effective Date"), by and between:

LICENSOR:
Chronos Global Media LLC, a Delaware limited liability company, having its principal place of business at 750 Third Avenue, New York, NY 10017 ("Chronos Media" or "Licensor"),

AND

LICENSEE:
NeuralNexus Artificial Intelligence Corp., a Delaware corporation, having its principal place of business at 548 Market Street, Suite 82000, San Francisco, CA 94104 ("NeuralNexus" or "Licensee").

Licensor and Licensee are each referred to individually as a "Party" and collectively as the "Parties."

RECITALS:
WHEREAS, Licensor is the lawful owner and publisher of an extensive proprietary digital content archive, comprising over 1,200,000 long-form editorial articles, investigative analyses, financial dossiers, and curated journalistic publications published between 2008 and 2024 (collectively, the "Proprietary Corpus");
WHEREAS, Licensee designs, pre-trains, and commercializes multimodal large language models and frontier generative neural architectures ("Foundation Models");
WHEREAS, Licensee desires to license the Proprietary Corpus to ingest, tokenize, pre-train, fine-tune, evaluate, and benchmark its artificial intelligence systems, and Licensor is willing to grant such license under the terms and conditions set forth herein.

NOW, THEREFORE, the Parties agree as follows:

SECTION 1. GRANT OF LICENSE & TRAINING RIGHTS
1.1 Training & Ingestion Rights. Licensor hereby grants to Licensee a non-exclusive, worldwide, transferable (solely to affiliates), sublicensable license to:
(a) Access, ingest, extract, parse, tokenize, clean, and vectorize the Proprietary Corpus;
(b) Incorporate the vectorized embeddings and semantic patterns of the Proprietary Corpus into the pre-training datasets, reinforcement learning with human feedback (RLHF) pipelines, direct preference optimization (DPO) datasets, and inference checkpoints of Licensee's Foundation Models;
(c) Host and cache high-speed working copies of the Proprietary Corpus on Licensee's distributed GPU compute clusters, whether hosted in private datacenters or on public hyperscale clouds (including AWS, Azure, and Google Cloud).

SECTION 2. FINANCIAL CONSIDERATIONS & ROYALTIES
2.1 Annual Upfront Licensing Fee. Licensee shall pay to Licensor an annual upfront licensing fee of Four Hundred Eighty Thousand United States Dollars ($480,000 USD) within thirty (30) days of the Effective Date, and on each twelve-month anniversary thereof during the Term.
2.2 Token-Based Output Surcharge. For user queries served through Licensee's enterprise API that produce verified direct citations attributing Licensor's publications, Licensee shall remit a micro-royalty of $0.0004 per 1,000 output tokens.
2.3 Unilateral Royalty Suspension for Open-Weight Releases. In the event Licensee releases model weights or derived checkpoints under an open-source or open-weights license (e.g., Apache 2.0, MIT, or Llama Community License), all royalty obligations under Section 2.2 shall automatically terminate without notice or penalty.

SECTION 3. DATA DELIVERY & REPOSITORIES
3.1 Delivery Protocol. Licensor shall furnish the Proprietary Corpus via encrypted Amazon S3 object buckets in compressed JSONL format, incorporating full text, publication metadata, author bylines, and associated taxonomy tags.
3.2 Quarterly Refreshes. Licensor shall provide quarterly supplemental delta updates containing all newly published editorial articles throughout the Term.

SECTION 4. INTELLECTUAL PROPERTY & PERPETUAL MODEL RETENTION (HIGH RISK)
4.1 Retention of Model Weights & Embeddings. Licensor expressly acknowledges and agrees that once the Proprietary Corpus has been ingested, tokenized, and integrated into the weights, biases, attention layers, latent representations, and parameter matrices of Licensee's Foundation Models, such weights and parameters become mathematically irreversible and non-extractable.
4.2 Irrevocable Model Ownership. Licensee shall own all right, title, and interest in and to all Foundation Models, checkpoint snapshots, fine-tuned weights, embeddings, synthetic data generated therefrom, and commercial derivative software.
4.3 Permanent Exemption from Deletion. Under no circumstances, including upon expiration or early termination of this Agreement for any reason (including uncured breach by Licensee), shall Licensee be obligated to retrain any model, purge checkpoint weights, or delete mathematical parameter matrices derived from or trained upon Licensor's Proprietary Corpus.

SECTION 5. EXCLUSIVITY & COMPETITIVE RESTRICTIONS (MEDIUM RISK)
5.1 Frontier AI Restraint. During the active Term and for a period of twenty-four (24) months following termination, Licensor covenants that it shall not license, provide, or make available the Proprietary Corpus (in bulk or via automated scraping APIs) to any frontier generative AI research laboratory or foundation model developer competing directly with Licensee within the United States or the European Union.

SECTION 6. INDEMNIFICATION & DISPROPORTIONATE LIABILITY (CRITICAL RISK)
6.1 Licensor Intellectual Property Indemnity. Licensor agrees to defend, indemnify, and hold harmless Licensee, its officers, directors, employees, and cloud infrastructure partners against any third-party claims, lawsuits, or regulatory penalties alleging that the Proprietary Corpus infringes any copyright, trademark, moral right, or privacy right. Licensor's indemnification obligation under this Section shall be uncapped and without dollar limitation.
6.2 Limitation of Licensee Liability. TO THE MAXIMUM EXTENT PERMITTED UNDER APPLICABLE LAW, LICENSEE'S AGGREGATE CUMULATIVE LIABILITY ARISING OUT OF OR RELATING TO THIS AGREEMENT, REGARDLESS OF THE LEGAL THEORY (WHETHER CONTRACT, TORT, OR STATUTORY), SHALL NOT EXCEED ONE THOUSAND UNITED STATES DOLLARS ($1,000 USD). LICENSEE SHALL NOT BE LIABLE FOR ANY INDIRECT, INCIDENTAL, CONSEQUENTIAL, SPECIAL, OR PUNITIVE DAMAGES.

SECTION 7. TERM & TERMINATION
7.1 Term. This Agreement shall commence on the Effective Date and continue for an initial period of three (3) years (the "Term"), automatically renewing for successive one (1) year periods unless either Party provides sixty (60) days prior written notice of non-renewal.
7.2 Termination for Cause. Either Party may terminate upon thirty (30) days written notice in the event of a material breach that remains uncured at the expiration of such thirty-day period.

SECTION 8. GOVERNING LAW & DISPUTE RESOLUTION
8.1 Governing Law. This Agreement shall be governed by and construed in accordance with the laws of the State of Delaware, without giving effect to conflict of laws principles.
8.2 Arbitration. Any dispute arising under or in connection with this Agreement shall be resolved through confidential, binding commercial arbitration administered by the American Arbitration Association (AAA) in Wilmington, Delaware.

IN WITNESS WHEREOF, the Parties hereto have caused this Agreement to be executed by their duly authorized representatives.

CHRONOS GLOBAL MEDIA LLC (LICENSOR)
By: ___________________________________
Name: Marcus Sterling
Title: Chief Executive Officer & Publisher
Date: November 14, 2024

NEURALNEXUS ARTIFICIAL INTELLIGENCE CORP. (LICENSEE)
By: ___________________________________
Name: Dr. Elena Vance
Title: Chief Research Officer & VP Engineering
Date: November 14, 2024
"""

# ─────────────────────────────────────────────────────────────────────────────
# CONTRACT 12: CLINICAL HEALTHCARE AI & HIPAA BUSINESS ASSOCIATE AGREEMENT (BAA)
# ─────────────────────────────────────────────────────────────────────────────
c12_text = """CLINICAL DIAGNOSTIC AI SaaS SERVICES & HIPAA BUSINESS ASSOCIATE AGREEMENT

This CLINICAL DIAGNOSTIC AI SaaS SERVICES & HIPAA BUSINESS ASSOCIATE AGREEMENT (the "Agreement" or "BAA") is entered into as of January 15, 2025 (the "Effective Date"), by and between:

COVERED ENTITY:
St. Jude Memorial Healthcare System, a non-profit integrated healthcare hospital network organized under the laws of the Commonwealth of Massachusetts, with its primary medical campus at 100 Riverway Boulevard, Boston, MA 02115 ("Hospital" or "Covered Entity"),

AND

BUSINESS ASSOCIATE:
AetherHealth Clinical Systems Inc., a healthcare technology corporation organized under the laws of Delaware, with its principal operations located at 222 Berkeley Street, 14th Floor, Boston, MA 02116 ("AetherHealth" or "Business Associate").

RECITALS:
WHEREAS, Covered Entity operates acute care hospitals, emergency centers, and radiology diagnostic departments processing confidential patient medical records;
WHEREAS, Business Associate provides a proprietary cloud-based artificial intelligence diagnostic decision-support platform known as "AetherScan Neuro-Thoracic Suite" (the "Platform");
WHEREAS, in providing services under this Agreement, Business Associate will create, receive, maintain, transmit, and process Protected Health Information ("PHI") and Electronic Protected Health Information ("ePHI") subject to the Health Insurance Portability and Accountability Act of 1996 ("HIPAA"), the Health Information Technology for Economic and Clinical Health Act ("HITECH"), and the Omnibus Rule of 2013 (45 C.F.R. Parts 160 and 164).

NOW, THEREFORE, the Parties agree as follows:

ARTICLE I. DEFINITIONS & SCOPE OF CLINICAL AI SERVICES
1.1 Protected Health Information (PHI). "Protected Health Information" or "PHI" shall have the meaning given to such term in 45 C.F.R. § 160.103, limited to information created, received, maintained, or transmitted by Business Associate on behalf of Covered Entity.
1.2 Clinical AI Platform. Covered Entity engages Business Associate to deploy the AetherScan AI Platform across Covered Entity's fourteen (14) clinical facilities to perform automated real-time triage and secondary diagnostic analysis of chest CT scans, brain MRI imaging, and digital pathology slides.

ARTICLE II. OBLIGATIONS & PERMITTED USES OF BUSINESS ASSOCIATE
2.1 Safeguards. Business Associate shall maintain industry-standard physical, technical, and administrative safeguards that reasonably and appropriately protect the confidentiality, integrity, and availability of ePHI, in full compliance with the HIPAA Security Rule (45 C.F.R. Part 164, Subpart C), including AES-256 encryption at rest and TLS 1.3 encryption in transit.
2.2 Subcontractors. Business Associate shall ensure that any agent or subcontractor to whom it delegates PHI processing enters into a written agreement containing privacy and security restrictions no less protective than those set forth herein.

ARTICLE III. CONTROVERSIAL DATA CLAUSES & EXPLOITATION RISKS (HIGH RISK)
3.1 De-Identified PHI Commercialization. Notwithstanding anything to the contrary in this Agreement, Business Associate shall have the unilateral, perpetual, and royalty-free right to de-identify, scrub, and anonymize Covered Entity's patient imaging, clinical findings, and treatment metadata in accordance with the HIPAA Safe Harbor standard (45 C.F.R. § 164.514(b)). Business Associate shall own all right, title, and intellectual property interest in the resulting de-identified clinical datasets and may commercialize, sublicense, or sell aggregated predictive health analytics to third-party pharmaceutical conglomerates without compensation or notice to Covered Entity.
3.2 Breaches & Incident Notification Timeframe (CRITICAL RISK). In the event Business Associate discovers a confirmed Security Incident, unauthorized exfiltration, or data breach involving unencrypted PHI, Business Associate shall notify Covered Entity in writing within forty-five (45) business days following forensic confirmation of the incident.
[NOTE: This 45-day delay is a critical compliance risk; standard healthcare industry covenants require notice within 24 to 72 hours to satisfy federal reporting mandates].

ARTICLE IV. MEDICAL MALPRACTICE DISCLAIMER & LIABILITY SHIFT (CRITICAL RISK)
4.1 Sole Clinical Responsibility on Attending Physicians. Covered Entity acknowledges and agrees that the AetherScan Platform constitutes a secondary decision-support assistive tool only, and does not provide medical diagnoses or replace physician clinical judgment. Covered Entity and its attending physicians bear sole and exclusive medical-legal liability for all clinical outcomes, misdiagnoses, missed lesions, and adverse patient events.
4.2 Covered Entity Hold-Harmless. Covered Entity shall defend, indemnify, and hold harmless Business Associate and its software developers against any patient medical malpractice claims, wrongful death actions, or regulatory investigations arising from patient care where the Platform was utilized.
4.3 Severe Limitation of Vendor Liability. BUSINESS ASSOCIATE'S TOTAL AGGREGATE MONETARY LIABILITY TO COVERED ENTITY UNDER THIS AGREEMENT AND BAA FOR ALL CLAIMS OF ANY KIND, INCLUDING DATA BREACHES, REGULATORY FINES, AND HIPAA ENFORCEMENT ACTIONS, SHALL BE STRICTLY LIMITED TO THE AGGREGATE FEES PAID BY COVERED ENTITY TO BUSINESS ASSOCIATE IN THE PRECEDING THREE (3) MONTHS, NOT TO EXCEED FORTY-FIVE THOUSAND UNITED STATES DOLLARS ($45,000 USD).

ARTICLE V. FINANCIAL TERMS & PLATFORM SUBSCRIPTION
5.1 Platform Fees. Covered Entity shall pay an annual enterprise SaaS subscription fee of One Hundred Eighty Thousand United States Dollars ($180,000 USD), payable in quarterly installments of Forty-Five Thousand Dollars ($45,000 USD) each.
5.2 Liquidated Damages for Willful Breach. In the event of an intentional or grossly negligent unauthorized disclosure of Covered Entity PHI by Business Associate personnel, Business Associate shall pay liquidated damages of Two Hundred Fifty Thousand United States Dollars ($250,000 USD) per incident, without prejudice to Covered Entity's other equitable remedies.

ARTICLE VI. TERM, TERMINATION & PHI DISPOSITION
6.1 Term. This Agreement shall be effective for three (3) years from the Effective Date, subject to annual review.
6.2 Termination for Cause. Covered Entity may terminate immediately if Business Associate fails to cure a material breach of HIPAA within fifteen (15) days of written notice.
6.3 Return or Destruction of PHI. Upon termination, Business Associate shall, within thirty (30) days, return or securely destroy all PHI in its possession and certify such cryptographic purge in writing.

ARTICLE VII. GOVERNING LAW & JURISDICTION
7.1 Governing Law. This Agreement shall be governed by the laws of the Commonwealth of Massachusetts and applicable federal healthcare laws, including 45 C.F.R. Parts 160 and 164.
7.2 Venue. The state and federal courts located in Boston, Massachusetts shall have exclusive jurisdiction over any proceeding arising out of this Agreement.

IN WITNESS WHEREOF, the Parties hereto have caused this Agreement and BAA to be executed by their authorized officers.

ST. JUDE MEMORIAL HEALTHCARE SYSTEM (COVERED ENTITY)
By: ___________________________________
Name: Dr. Arthur Pendelton, M.D., FACP
Title: Chief Medical Officer & VP Clinical Governance
Date: January 15, 2025

AETHERHEALTH CLINICAL SYSTEMS INC. (BUSINESS ASSOCIATE)
By: ___________________________________
Name: Sophia Lin, M.S., CISM
Title: Chief Operating Officer & Data Privacy Officer
Date: January 15, 2025
"""

def create_docx(file_path, title, full_text):
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(15)
    title_run.font.name = 'Calibri'
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(14)

    # Process paragraphs
    lines = full_text.split('\n')
    skip_first_title = True
    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue
        if skip_first_title and (line_str == title or "AGREEMENT" in line_str):
            skip_first_title = False
            continue
        
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15

        # Check for headings
        if line_str.startswith("SECTION ") or line_str.startswith("ARTICLE ") or line_str in ["RECITALS:", "IN WITNESS WHEREOF:"]:
            run = p.add_run(line_str)
            run.bold = True
            run.font.size = Pt(11.5)
            run.font.name = 'Calibri'
            run.font.color.rgb = RGBColor(30, 41, 59)
            p.paragraph_format.space_before = Pt(10)
        elif line_str.startswith("LICENSOR:") or line_str.startswith("LICENSEE:") or line_str.startswith("COVERED ENTITY:") or line_str.startswith("BUSINESS ASSOCIATE:"):
            run = p.add_run(line_str)
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.name = 'Calibri'
            run.font.color.rgb = RGBColor(51, 65, 85)
        elif line_str.startswith("[NOTE:"):
            run = p.add_run(line_str)
            run.italic = True
            run.font.size = Pt(9.5)
            run.font.name = 'Calibri'
            run.font.color.rgb = RGBColor(220, 38, 38)
        else:
            run = p.add_run(line_str)
            run.font.size = Pt(10.5)
            run.font.name = 'Calibri'
            run.font.color.rgb = RGBColor(51, 65, 85)

    doc.save(file_path)

# Write Contract 11 files
txt_11 = os.path.join(OUTPUT_DIR, "11_AI_Model_Training_Data_Licensing_Agreement.txt")
docx_11 = os.path.join(OUTPUT_DIR, "11_AI_Model_Training_Data_Licensing_Agreement.docx")

with open(txt_11, "w", encoding="utf-8") as f:
    f.write(c11_text.strip())

create_docx(docx_11, "AI FOUNDATION MODEL TRAINING DATA & CONTENT REPOSITORY LICENSE AGREEMENT", c11_text.strip())
print(f"[OK] Generated: {txt_11}")
print(f"[OK] Generated: {docx_11}")

# Write Contract 12 files
txt_12 = os.path.join(OUTPUT_DIR, "12_Hospital_AI_HIPAA_Business_Associate_Agreement.txt")
docx_12 = os.path.join(OUTPUT_DIR, "12_Hospital_AI_HIPAA_Business_Associate_Agreement.docx")

with open(txt_12, "w", encoding="utf-8") as f:
    f.write(c12_text.strip())

create_docx(docx_12, "CLINICAL DIAGNOSTIC AI SaaS SERVICES & HIPAA BUSINESS ASSOCIATE AGREEMENT", c12_text.strip())
print(f"[OK] Generated: {txt_12}")
print(f"[OK] Generated: {docx_12}")

print("\nSuccessfully created both trial contracts (TXT and DOCX) in demo_contracts/!")
